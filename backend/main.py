import os
import shutil
import json
import time
import re
import math
import logging
from datetime import datetime
from typing import Optional
from fastapi import FastAPI, UploadFile, File, Depends, HTTPException, Form, Query, BackgroundTasks
from fastapi.security import OAuth2PasswordRequestForm
from pydantic import BaseModel
from fastapi.responses import FileResponse
from fastapi.middleware.cors import CORSMiddleware
from sqlalchemy.orm import Session
from sqlalchemy import func
from app.database import get_db, engine, Base
from app.models import Document, DocumentChunk, Folder, User, BlacklistedToken, Vacancy
from app.auth import get_current_user, create_access_token, verify_password, oauth2_scheme, require_role
from app.crypto import encrypt_sensitive_data, decrypt_sensitive_data, mask_secret, sanitize_error_log
from app.notifications import notify_document_status
import google.generativeai as genai
import pdfplumber
import docx
from dotenv import load_dotenv

load_dotenv()
logger = logging.getLogger("IndicioBackend")

# Configuración y enmascaramiento seguro de IA
gemini_key = os.getenv("GEMINI_API_KEY")
if gemini_key:
    logger.info("API Key de Gemini configurada de forma segura: %s", mask_secret(gemini_key))
    genai.configure(api_key=gemini_key)
else:
    logger.warning("[SEGURIDAD CRÍTICA] GEMINI_API_KEY no encontrada en las variables de entorno.")

MODEL_TEXT = os.getenv("GEMINI_TEXT_MODEL", "gemini-2.0-flash")
MODEL_EMB = os.getenv("GEMINI_EMBEDDING_MODEL", "models/text-embedding-005")

app = FastAPI(title="Indicio Enterprise", description="Plataforma Segura de Gestión Documental Inteligente y RAG")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"], 
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

from sqlalchemy import text

@app.on_event("startup")
def startup_db_check():
    """Garantiza la consistencia automática del esquema de la base de datos al encender el backend"""
    try:
        with engine.connect() as conn:
            conn.execute(text("CREATE EXTENSION IF NOT EXISTS vector;"))
            conn.execute(text("ALTER TABLE users ADD COLUMN IF NOT EXISTS role VARCHAR(50) DEFAULT 'user';"))
            conn.execute(text("ALTER TABLE users ADD COLUMN IF NOT EXISTS tenant_id INTEGER REFERENCES tenants(id) ON DELETE CASCADE;"))
            conn.commit()
        logger.info("[SEGURIDAD] Verificación y migración de esquema completada exitosamente.")
    except Exception as e:
        logger.warning("[SEGURIDAD] No se pudo verificar el esquema en startup: %s", e)

def serialize_document(doc: Document, user_role: str) -> dict:
    """
    Serializa un documento aplicando:
    1. Descifrado seguro de datos sensibles en reposo (_encrypted_data).
    2. Restricción de acceso a registros técnicos / error_log según el rol del usuario.
    3. Sanitización de trazas para garantizar que ninguna API key o token sea devuelto.
    """
    data = doc.extracted_data
    if isinstance(data, dict) and "_encrypted_data" in data:
        decrypted = decrypt_sensitive_data(data["_encrypted_data"])
        if isinstance(decrypted, dict):
            data = decrypted

    # Restricción de registros: Solo administradores y auditores tienen acceso al detalle técnico
    error_log = None
    if user_role in ["admin", "auditor"]:
        error_log = sanitize_error_log(doc.error_log)
    elif doc.status == "Error":
        error_log = "Error en el procesamiento del archivo. Contacte al administrador del sistema."

    return {
        "id": doc.id,
        "filename": doc.filename,
        "category": doc.category,
        "status": doc.status,
        "folder_id": doc.folder_id,
        "vacancy_id": doc.vacancy_id,
        "mime_type": doc.mime_type,
        "file_size_bytes": doc.file_size_bytes,
        "summary": doc.summary,
        "extracted_data": data,
        "error_log": error_log,
        "created_at": doc.created_at,
        "tenant_id": doc.tenant_id
    }

# --- SERVICIOS ---
def extract_text(file_path: str, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    text_content = ""
    if ext == ".pdf":
        with pdfplumber.open(file_path) as pdf:
            text_content = "\n".join([page.extract_text() for page in pdf.pages if page.extract_text()])
    elif ext == ".docx":
        doc = docx.Document(file_path)
        text_content = "\n".join([para.text for para in doc.paragraphs])
    elif ext == ".txt":
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text_content = f.read()
    return text_content.strip()

# --- ENDPOINTS ---
class LoginRequest(BaseModel):
    email: str
    password: str

@app.post("/api/auth/login")
def login(login_data: LoginRequest, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.email == login_data.email).first()
    if not user or not verify_password(login_data.password, user.hashed_password):
        raise HTTPException(status_code=400, detail="Email o contraseña incorrectos")
    
    user_role = getattr(user, "role", "user") or "user"
    tenant_id = user.tenant_id if user.tenant_id is not None else 1
    access_token = create_access_token(data={"sub": user.email, "role": user_role, "tenant_id": tenant_id})
    return {
        "access_token": access_token, 
        "token_type": "bearer",
        "role": user_role,
        "tenant_id": tenant_id
    }

@app.post("/api/auth/logout")
def logout(token: str = Depends(oauth2_scheme), db: Session = Depends(get_db)):
    blacklisted = BlacklistedToken(token=token)
    db.add(blacklisted)
    db.commit()
    return {"message": "Sesión cerrada exitosamente"}

@app.get("/api/dashboard/stats")
def get_stats(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    total = db.query(Document).filter(Document.tenant_id == current_user.tenant_id).count()
    processed = db.query(Document).filter(Document.tenant_id == current_user.tenant_id, Document.status == "Procesado").count()
    errors = db.query(Document).filter(Document.tenant_id == current_user.tenant_id, Document.status == "Error").count()
    
    categories_query = db.query(Document.category, func.count(Document.id)).filter(Document.tenant_id == current_user.tenant_id).group_by(Document.category).all()
    categories = {cat if cat else "Sin Clasificar": count for cat, count in categories_query}
    
    return {
        "total": total, 
        "processed": processed, 
        "errors": errors,
        "categories": categories
    }

class FolderCreate(BaseModel):
    name: str

@app.post("/api/folders")
def create_folder(folder: FolderCreate, db: Session = Depends(get_db), current_user: User = Depends(require_role(["admin"]))):
    new_folder = Folder(name=folder.name, tenant_id=current_user.tenant_id)
    db.add(new_folder)
    db.commit()
    db.refresh(new_folder)
    return new_folder

@app.get("/api/folders")
def list_folders(db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    return db.query(Folder).filter(Folder.tenant_id == current_user.tenant_id).order_by(Folder.created_at.desc()).all()

@app.delete("/api/folders/{folder_id}")
def delete_folder(folder_id: int, db: Session = Depends(get_db), current_user: User = Depends(require_role(["admin"]))):
    folder = db.query(Folder).filter(Folder.id == folder_id, Folder.tenant_id == current_user.tenant_id).first()
    if not folder:
        raise HTTPException(status_code=404, detail="Carpeta no encontrada o no pertenece a su organización")
    db.delete(folder)
    db.commit()
    return {"message": "Carpeta eliminada"}

# --- VACANTES Y EMPAREJAMIENTO HÍBRIDO (NIVEL 3) ---

def cosine_similarity(vec1, vec2) -> float:
    if not vec1 or not vec2:
        return 0.0
    v1 = list(vec1)
    v2 = list(vec2)
    if len(v1) != len(v2):
        return 0.0
    dot = sum(a * b for a, b in zip(v1, v2))
    norm1 = math.sqrt(sum(a * a for a in v1))
    norm2 = math.sqrt(sum(b * b for b in v2))
    if norm1 == 0.0 or norm2 == 0.0:
        return 0.0
    return dot / (norm1 * norm2)

def compute_vacancy_embedding(title: str, department: str, description: str, requirements: str, required_skills: list) -> list:
    skills_str = ", ".join(required_skills) if isinstance(required_skills, list) else str(required_skills or "")
    full_text = f"Vacante: {title}. Departamento: {department}. Descripción del puesto: {description}. Requisitos: {requirements}. Habilidades requeridas: {skills_str}"
    try:
        res = genai.embed_content(
            model=MODEL_EMB,
            content=full_text[:2000],
            task_type="retrieval_document",
            output_dimensionality=768
        )
        return res['embedding']
    except Exception as e:
        print(f"[WARN] Error generando embedding de vacante: {e}")
        return None

def calculate_semantic_similarity(vacancy_emb, doc_chunks: list) -> float:
    if not vacancy_emb or not doc_chunks:
        return 70.0
    similarities = []
    for chunk in doc_chunks:
        if chunk.embedding is not None:
            sim = cosine_similarity(vacancy_emb, chunk.embedding)
            similarities.append(sim)
    if not similarities:
        return 70.0
    similarities.sort(reverse=True)
    top_sim = similarities[:3]
    avg_sim = sum(top_sim) / len(top_sim)
    score = max(0.0, min(100.0, ((avg_sim + 1.0) / 2.0) * 100.0))
    return round(score, 1)

def evaluate_cv_against_vacancy(cv_text: str, vacancy: Vacancy, candidate_models: list) -> dict:
    skills_str = ", ".join(vacancy.required_skills) if isinstance(vacancy.required_skills, list) else str(vacancy.required_skills or "")
    prompt = f"""Eres un Headhunter Senior y evaluador técnico de talento humano.
Evalúa ESTRICTAMENTE este currículum contra la VACANTE OBJETIVO indicada.
Tu análisis debe ser objetivo, técnico y cuantitativo.
Responde ÚNICAMENTE un objeto JSON válido, sin delimitadores markdown.

VACANTE OBJETIVO:
- Título del Puesto: {vacancy.title}
- Departamento: {vacancy.department}
- Descripción: {vacancy.description}
- Requisitos: {vacancy.requirements}
- Años mínimos de experiencia requeridos: {vacancy.min_years_experience}
- Habilidades / Tecnologías clave requeridas: {skills_str}

CURRÍCULUM DEL POSTULANTE:
{cv_text[:7000]}

ESQUEMA JSON REQUERIDO:
{{
    "postulante": "Nombre completo detectado",
    "rol": "Cargo o perfil profesional principal",
    "anos_experiencia": 0.0,
    "competencias": ["habilidad1", "habilidad2"],
    "matched_skills": ["habilidad requerida por la vacante que el candidato SÍ posee"],
    "missing_skills": ["habilidad requerida por la vacante que el candidato NO menciona o carece"],
    "experience_gap": "Resumen del cumplimiento de experiencia (ej: 'Cumple: 4 años vs 3 mínimos')",
    "green_flags": ["Fortaleza 1 frente a la vacante", "Fortaleza 2"],
    "red_flags": ["Brecha o riesgo frente a la vacante"],
    "gemini_score": 85.0,
    "resumen_scouting": "Síntesis ejecutiva de 2 a 3 líneas sobre el ajuste del candidato a esta vacante"
}}"""

    for model_name in candidate_models:
        try:
            curr_model = genai.GenerativeModel(model_name)
            res = curr_model.generate_content(prompt, generation_config={"response_mime_type": "application/json"})
            raw = res.text.strip()
            if raw.startswith("```json"):
                raw = raw[7:-3].strip()
            elif raw.startswith("```"):
                raw = raw[3:-3].strip()
            return json.loads(raw)
        except Exception as e:
            print(f"[WARN] Error en evaluate_cv_against_vacancy con {model_name}: {e}")
            continue
    return None

class VacancyCreate(BaseModel):
    title: str
    department: Optional[str] = "Tecnología"
    description: str
    requirements: str
    min_years_experience: Optional[int] = 0
    required_skills: Optional[list[str]] = []
    status: Optional[str] = "Abierta"

class VacancyUpdate(BaseModel):
    title: Optional[str] = None
    department: Optional[str] = None
    description: Optional[str] = None
    requirements: Optional[str] = None
    min_years_experience: Optional[int] = None
    required_skills: Optional[list[str]] = None
    status: Optional[str] = None

@app.post("/api/vacancies")
def create_vacancy(
    data: VacancyCreate, 
    db: Session = Depends(get_db), 
    current_user: User = Depends(require_role(["admin"]))
):
    emb = compute_vacancy_embedding(data.title, data.department, data.description, data.requirements, data.required_skills)
    new_vac = Vacancy(
        tenant_id=current_user.tenant_id,
        title=data.title,
        department=data.department or "Tecnología",
        description=data.description,
        requirements=data.requirements,
        min_years_experience=data.min_years_experience or 0,
        required_skills=data.required_skills or [],
        status=data.status or "Abierta",
        embedding=emb
    )
    db.add(new_vac)
    db.commit()
    db.refresh(new_vac)
    return new_vac

@app.get("/api/vacancies")
def list_vacancies(
    status: Optional[str] = Query(None),
    db: Session = Depends(get_db), 
    current_user: User = Depends(get_current_user)
):
    q = db.query(Vacancy).filter(Vacancy.tenant_id == current_user.tenant_id)
    if status:
        q = q.filter(Vacancy.status == status)
    vacancies = q.order_by(Vacancy.created_at.desc()).all()
    
    result = []
    for vac in vacancies:
        candidates = db.query(Document).filter(Document.vacancy_id == vac.id, Document.tenant_id == current_user.tenant_id).all()
        scores = []
        for c in candidates:
            if c.extracted_data and isinstance(c.extracted_data, dict):
                score = c.extracted_data.get("match_score")
                if score is not None:
                    try:
                        scores.append(float(score))
                    except Exception:
                        pass
        top_score = max(scores) if scores else 0.0
        avg_score = round(sum(scores) / len(scores), 1) if scores else 0.0
        
        result.append({
            "id": vac.id,
            "title": vac.title,
            "department": vac.department,
            "description": vac.description,
            "requirements": vac.requirements,
            "min_years_experience": vac.min_years_experience,
            "required_skills": vac.required_skills or [],
            "status": vac.status,
            "created_at": vac.created_at,
            "candidates_count": len(candidates),
            "top_score": top_score,
            "avg_score": avg_score
        })
    return result

@app.get("/api/vacancies/{vacancy_id}")
def get_vacancy_detail(
    vacancy_id: int, 
    db: Session = Depends(get_db), 
    current_user: User = Depends(get_current_user)
):
    vac = db.query(Vacancy).filter(Vacancy.id == vacancy_id, Vacancy.tenant_id == current_user.tenant_id).first()
    if not vac:
        raise HTTPException(status_code=404, detail="Vacante no encontrada")
        
    candidates = db.query(Document).filter(Document.vacancy_id == vac.id, Document.tenant_id == current_user.tenant_id).all()
    
    def get_score(doc):
        if doc.extracted_data and isinstance(doc.extracted_data, dict):
            s = doc.extracted_data.get("match_score")
            if s is not None:
                try:
                    return float(s)
                except Exception:
                    pass
        return 0.0
        
    sorted_candidates = sorted(candidates, key=get_score, reverse=True)
    
    return {
        "id": vac.id,
        "title": vac.title,
        "department": vac.department,
        "description": vac.description,
        "requirements": vac.requirements,
        "min_years_experience": vac.min_years_experience,
        "required_skills": vac.required_skills or [],
        "status": vac.status,
        "created_at": vac.created_at,
        "candidates": sorted_candidates
    }

@app.put("/api/vacancies/{vacancy_id}")
def update_vacancy(
    vacancy_id: int,
    data: VacancyUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role(["admin"]))
):
    vac = db.query(Vacancy).filter(Vacancy.id == vacancy_id, Vacancy.tenant_id == current_user.tenant_id).first()
    if not vac:
        raise HTTPException(status_code=404, detail="Vacante no encontrada")
        
    need_emb = False
    if data.title is not None:
        vac.title = data.title
        need_emb = True
    if data.department is not None:
        vac.department = data.department
    if data.description is not None:
        vac.description = data.description
        need_emb = True
    if data.requirements is not None:
        vac.requirements = data.requirements
        need_emb = True
    if data.min_years_experience is not None:
        vac.min_years_experience = data.min_years_experience
    if data.required_skills is not None:
        vac.required_skills = data.required_skills
        need_emb = True
    if data.status is not None:
        vac.status = data.status
        
    if need_emb:
        vac.embedding = compute_vacancy_embedding(vac.title, vac.department, vac.description, vac.requirements, vac.required_skills)
        
    db.commit()
    db.refresh(vac)
    return vac

@app.delete("/api/vacancies/{vacancy_id}")
def delete_vacancy(
    vacancy_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role(["admin"]))
):
    vac = db.query(Vacancy).filter(Vacancy.id == vacancy_id, Vacancy.tenant_id == current_user.tenant_id).first()
    if not vac:
        raise HTTPException(status_code=404, detail="Vacante no encontrada")
    db.delete(vac)
    db.commit()
    return {"message": "Vacante eliminada"}

@app.post("/api/vacancies/{vacancy_id}/match/{document_id}")
def match_document_to_vacancy(
    vacancy_id: int,
    document_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role(["admin"]))
):
    from sqlalchemy.orm.attributes import flag_modified
    vac = db.query(Vacancy).filter(Vacancy.id == vacancy_id, Vacancy.tenant_id == current_user.tenant_id).first()
    if not vac:
        raise HTTPException(status_code=404, detail="Vacante no encontrada")
    
    doc = db.query(Document).filter(Document.id == document_id, Document.tenant_id == current_user.tenant_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Documento no encontrado")
        
    raw_text = extract_text(doc.file_path, doc.filename)
    if not raw_text:
        raise HTTPException(status_code=400, detail="No se pudo extraer texto del documento")
        
    chunks = db.query(DocumentChunk).filter(DocumentChunk.document_id == doc.id).all()
    semantic_score = calculate_semantic_similarity(vac.embedding, chunks)
    
    candidate_models = list(dict.fromkeys([
        os.getenv("GEMINI_TEXT_MODEL", "gemini-3.8-flash"),
        "gemini-3.8-flash",
        "gemini-3.7-flash",
        "gemini-3.5-flash-lite",
        "gemini-3.1-flash-lite"
    ]))
    vac_eval = evaluate_cv_against_vacancy(raw_text, vac, candidate_models)
    
    extracted = dict(doc.extracted_data or {})
    if vac_eval:
        gemini_score = float(vac_eval.get("gemini_score", 75.0))
        final_match_score = round(0.65 * gemini_score + 0.35 * semantic_score, 1)
        extracted["match_score"] = final_match_score
        extracted["gemini_score"] = gemini_score
        extracted["semantic_score"] = semantic_score
        extracted["vacancy_id"] = vac.id
        extracted["vacancy_title"] = vac.title
        extracted["matched_skills"] = vac_eval.get("matched_skills", [])
        extracted["missing_skills"] = vac_eval.get("missing_skills", [])
        extracted["experience_gap"] = vac_eval.get("experience_gap", "")
        extracted["green_flags"] = vac_eval.get("green_flags", extracted.get("green_flags", []))
        extracted["red_flags"] = vac_eval.get("red_flags", [])
        extracted["resumen_scouting"] = vac_eval.get("resumen_scouting", extracted.get("resumen_scouting", ""))
        if vac_eval.get("postulante"):
            extracted["postulante"] = vac_eval["postulante"]
        if vac_eval.get("rol"):
            extracted["rol"] = vac_eval["rol"]
        if vac_eval.get("anos_experiencia") is not None:
            try:
                extracted["anos_experiencia"] = float(vac_eval["anos_experiencia"])
            except Exception:
                pass
    else:
        extracted["match_score"] = semantic_score
        extracted["semantic_score"] = semantic_score
        extracted["vacancy_id"] = vac.id
        extracted["vacancy_title"] = vac.title

    doc.vacancy_id = vac.id
    doc.extracted_data = extracted
    flag_modified(doc, "extracted_data")
    db.commit()
    db.refresh(doc)
    
    return {
        "message": "Evaluación contra vacante completada",
        "document_id": doc.id,
        "vacancy_id": vac.id,
        "match_score": extracted.get("match_score"),
        "gemini_score": extracted.get("gemini_score"),
        "semantic_score": extracted.get("semantic_score"),
        "matched_skills": extracted.get("matched_skills"),
        "missing_skills": extracted.get("missing_skills"),
        "resumen_scouting": extracted.get("resumen_scouting")
    }

def process_document_background(doc_id: int, file_path: str, filename: str):
    # This must create its own DB session since it runs in the background
    db = next(get_db())
    doc = db.query(Document).filter(Document.id == doc_id).first()
    if not doc:
        return
        
    try:
        # Transición a Procesando antes de la IA
        doc.status = "Procesando"
        db.commit()
        
        # 1. Extracción con validación
        raw_text = extract_text(file_path, filename)
        if not raw_text or len(raw_text.strip()) < 10:
            raise ValueError("El documento está vacío o no se pudo extraer texto legible.")
        
        # 2. Clasificación con modelo dinámico y fallback ante límites de cuota (429)
        prompt = f"""Analiza este documento y responde ESTRICTAMENTE en formato JSON válido, sin comillas invertidas ni markdown.
        Clasifica obligatoriamente en una de estas categorías: "Facturas y Cuentas de Cobro", "Contratos y Acuerdos Legales", "Hojas de Vida / Perfiles Laborales", "Recibos y Comprobantes de Pago", "Reportes y Estados Financieros", "Actas de Reunión", "Manuales y Procedimientos" (o "Sin Clasificar" si no aplica ninguna).
        El esquema debe ser:
        {{
            "category": "nombre de la categoria",
            "summary": "resumen ejecutivo",
            "extracted_data": {{
                // Si es "Facturas y Cuentas de Cobro", usa OBLIGATORIAMENTE este esquema:
                // "emisor_nombre": string (nombre del proveedor o emisor),
                // "emisor_nit": string (NIT o identificación fiscal, o null si no aparece),
                // "numero_factura": string (número o consecutivo de factura),
                // "fecha_emision": string (YYYY-MM-DD o formato original),
                // "fecha_vencimiento": string (YYYY-MM-DD o null si no especifica),
                // "moneda": string (ej: "COP", "USD", "EUR"),
                // "monto_subtotal": number (numérico limpio float/int, sin símbolos $ ni puntos de miles, ej: 1200000.00),
                // "monto_impuestos": number (numérico limpio float/int, sin símbolos $ ni puntos de miles, ej: 228000.00),
                // "monto_total": number (numérico limpio float/int, sin símbolos $ ni puntos de miles, ej: 1428000.00),
                // "concepto_principal": string (descripción breve del servicio o bien)

                // Si es "Recibos y Comprobantes de Pago":
                // "emisor_nombre": string,
                // "receptor_nombre": string,
                // "numero_recibo": string (o null),
                // "fecha_pago": string,
                // "moneda": string,
                // "monto_total": number (numérico limpio float/int, sin símbolos ni puntos de miles),
                // "concepto_principal": string,
                // "metodo_pago": string (ej: "Transferencia", "Efectivo", o null)

                // Si es "Contratos y Acuerdos Legales": "partes", "objeto", "vigencia", "valor".
                // Si es "Hojas de Vida / Perfiles Laborales", usa OBLIGATORIAMENTE este esquema estructurado:
                // "postulante": string (nombre completo del candidato),
                // "rol": string (cargo o especialidad principal, ej: "Desarrollador Fullstack"),
                // "anos_experiencia": number (años de experiencia laboral corporativa en número, ej: 4 o 1.5),
                // "competencias": array de strings (tecnologías y habilidades clave, ej: ["React", "Azure", "GitHub", "Python"]),
                // "formacion": string (educación, certificaciones o nivel académico),
                // "contacto_telefono": string o null,
                // "contacto_ubicacion": string o null (ciudad o país),
                // "match_score": number (puntuación entera o decimal del 0 al 100 evaluando seniority, experiencia y solidez técnica),
                // "green_flags": array de strings (principales 2 a 3 fortalezas del candidato),
                // "resumen_scouting": string (síntesis ejecutiva de 2 líneas destacando el potencial del candidato)

                // Si es "Reportes y Estados Financieros": "empresa", "periodo", "tipo_reporte", "ingresos", "gastos".
                // Si es "Actas de Reunión": "fecha", "asistentes", "temas_tratados", "conclusiones".
                // Si es "Manuales y Procedimientos": "titulo", "departamento", "proposito", "version".
            }}
        }}

        REGLA ESTRICTA PARA MONTOS: Los campos 'monto_total', 'monto_subtotal' y 'monto_impuestos' deben ser EXCLUSIVAMENTE de tipo number (flotante o entero), NUNCA cadenas de texto, sin símbolos como '$', sin 'COP'/'USD' y sin separadores de miles con puntos. Ejemplo correcto: 1428000.00. Ejemplo incorrecto: "$ 1.428.000".

        Documento: {raw_text[:8000]}"""

        candidate_models = list(dict.fromkeys([
            os.getenv("GEMINI_TEXT_MODEL", "gemini-3.8-flash"),
            "gemini-3.8-flash",
            "gemini-3.7-flash",
            "gemini-3.5-flash-lite",
            "gemini-3.1-flash-lite"
        ]))
        
        response = None
        last_error = None
        for model_name in candidate_models:
            try:
                curr_model = genai.GenerativeModel(model_name)
                response = curr_model.generate_content(prompt, generation_config={"response_mime_type": "application/json"})
                print(f"[OK] Documento analizado exitosamente con {model_name}")
                break
            except Exception as e:
                last_error = e
                err_str = str(e)
                if "429" in err_str or "Quota exceeded" in err_str:
                    print(f"[WARN] Limite de cuota alcanzado en {model_name}. Cambiando a modelo alternativo...")
                    continue
                else:
                    print(f"[WARN] Error en {model_name}: {err_str[:80]}. Probando alternativo...")
                    continue

        if not response:
            raise last_error or ValueError("No se pudo obtener respuesta de la IA.")
            
        raw_response = response.text.strip()
        
        # Limpieza por si Gemini insiste en enviar Markdown
        if raw_response.startswith("```json"):
            raw_response = raw_response[7:-3].strip()
        elif raw_response.startswith("```"):
            raw_response = raw_response[3:-3].strip()
            
        ai_data = json.loads(raw_response)
        
        doc.category = ai_data.get("category", "Sin Clasificar")
        doc.summary = ai_data.get("summary", "")
        extracted = ai_data.get("extracted_data", {})
        
        # Sanitización defensiva de datos extraídos para garantizar compatibilidad total con n8n
        if isinstance(extracted, dict):
            for k in ["monto_total", "monto_subtotal", "monto_impuestos"]:
                if k in extracted and extracted[k] is not None:
                    if isinstance(extracted[k], str):
                        clean_str = extracted[k].replace("$", "").replace("USD", "").replace("COP", "").replace("EUR", "").replace(" ", "").strip()
                        try:
                            if "," in clean_str and "." in clean_str:
                                if clean_str.rfind(",") > clean_str.rfind("."):
                                    clean_str = clean_str.replace(".", "").replace(",", ".")
                                else:
                                    clean_str = clean_str.replace(",", "")
                            elif "," in clean_str:
                                clean_str = clean_str.replace(",", ".")
                            extracted[k] = float(clean_str)
                        except (ValueError, TypeError):
                            pass

            if "match_score" in extracted and extracted["match_score"] is not None:
                try:
                    if isinstance(extracted["match_score"], str):
                        clean_score = re.sub(r"[^\d.]", "", extracted["match_score"])
                        extracted["match_score"] = float(clean_score) if clean_score else 80.0
                    else:
                        extracted["match_score"] = float(extracted["match_score"])
                except Exception:
                    extracted["match_score"] = 80.0

            if "anos_experiencia" in extracted and extracted["anos_experiencia"] is not None:
                try:
                    if isinstance(extracted["anos_experiencia"], str):
                        clean_exp = re.sub(r"[^\d.]", "", extracted["anos_experiencia"])
                        extracted["anos_experiencia"] = float(clean_exp) if clean_exp else 0.0
                    else:
                        extracted["anos_experiencia"] = float(extracted["anos_experiencia"])
                except Exception:
                    pass
        doc.extracted_data = extracted

        # 3. Vectorización de Chunks
        chunks = [raw_text[i:i+700] for i in range(0, len(raw_text), 700)]
        created_chunks = []
        for i, chunk in enumerate(chunks):
            emb = None
            for attempt in range(3):
                try:
                    emb = genai.embed_content(model=MODEL_EMB, content=chunk, task_type="retrieval_document", output_dimensionality=768)['embedding']
                    break
                except Exception as e:
                    if "429" in str(e) or "Quota exceeded" in str(e):
                        if attempt < 2:
                            print(f"[WARN] Cuota excedida en embeddings. Esperando 20s... ({attempt+1}/3)")
                            time.sleep(20)
                            continue
                    raise e
            
            if emb:
                chunk_obj = DocumentChunk(document_id=doc.id, chunk_index=i, chunk_text=chunk, embedding=emb)
                db.add(chunk_obj)
                created_chunks.append(chunk_obj)

        # 4. Emparejamiento Híbrido Multivacante (Nivel 3) si es Hoja de Vida
        is_cv = doc.category and any(k in doc.category for k in ["Vida", "Curriculum", "CV", "Laboral"])
        target_vacancy = None
        if doc.vacancy_id:
            target_vacancy = db.query(Vacancy).filter(Vacancy.id == doc.vacancy_id).first()
        elif is_cv:
            open_vacancies = db.query(Vacancy).filter(Vacancy.tenant_id == doc.tenant_id, Vacancy.status == "Abierta").all()
            if open_vacancies:
                if len(open_vacancies) == 1:
                    target_vacancy = open_vacancies[0]
                else:
                    # Enrutamiento inteligente multivacante por afinidad semántica vectorial
                    best_vac = open_vacancies[0]
                    best_sim = -1.0
                    for vac in open_vacancies:
                        sim = calculate_semantic_similarity(vac.embedding, created_chunks)
                        if sim > best_sim:
                            best_sim = sim
                            best_vac = vac
                    target_vacancy = best_vac
                    print(f"[INFO] CV enrutado automáticamente a la vacante más afín: '{target_vacancy.title}' (Afinidad semántica: {best_sim}%)")
                doc.vacancy_id = target_vacancy.id

        if is_cv and target_vacancy:
            print(f"[INFO] Evaluando emparejamiento híbrido con vacante: {target_vacancy.title}")
            semantic_score = calculate_semantic_similarity(target_vacancy.embedding, created_chunks)
            vac_eval = evaluate_cv_against_vacancy(raw_text, target_vacancy, candidate_models)
            
            if vac_eval:
                gemini_score = float(vac_eval.get("gemini_score", 75.0))
                final_match_score = round(0.65 * gemini_score + 0.35 * semantic_score, 1)
                
                extracted["match_score"] = final_match_score
                extracted["gemini_score"] = gemini_score
                extracted["semantic_score"] = semantic_score
                extracted["vacancy_id"] = target_vacancy.id
                extracted["vacancy_title"] = target_vacancy.title
                extracted["matched_skills"] = vac_eval.get("matched_skills", [])
                extracted["missing_skills"] = vac_eval.get("missing_skills", [])
                extracted["experience_gap"] = vac_eval.get("experience_gap", "")
                extracted["green_flags"] = vac_eval.get("green_flags", extracted.get("green_flags", []))
                extracted["red_flags"] = vac_eval.get("red_flags", [])
                extracted["resumen_scouting"] = vac_eval.get("resumen_scouting", extracted.get("resumen_scouting", ""))
                if vac_eval.get("postulante"):
                    extracted["postulante"] = vac_eval["postulante"]
                if vac_eval.get("rol"):
                    extracted["rol"] = vac_eval["rol"]
                if vac_eval.get("anos_experiencia") is not None:
                    try:
                        extracted["anos_experiencia"] = float(vac_eval["anos_experiencia"])
                    except Exception:
                        pass
            else:
                extracted["match_score"] = semantic_score
                extracted["semantic_score"] = semantic_score
                extracted["vacancy_id"] = target_vacancy.id
                extracted["vacancy_title"] = target_vacancy.title
            
        # Cifrado de datos sensibles en PostgreSQL (Cifrado en reposo)
        doc.extracted_data = {
            "_encrypted_data": encrypt_sensitive_data(extracted),
            "match_score": extracted.get("match_score") if isinstance(extracted, dict) else None,
            "postulante": extracted.get("postulante") if isinstance(extracted, dict) else None,
            "rol": extracted.get("rol") if isinstance(extracted, dict) else None,
            "anos_experiencia": extracted.get("anos_experiencia") if isinstance(extracted, dict) else None,
            "total": extracted.get("total") if isinstance(extracted, dict) else None,
            "proveedor": extracted.get("proveedor") if isinstance(extracted, dict) else None
        }

        doc.status = "Procesado"
        print(f"[OK] Exito procesando: {filename}")

    except Exception as e:
        doc.status = "Error"
        doc.error_log = sanitize_error_log(str(e))
        print(f"[ERROR] Error en {filename}: {doc.error_log}") 
    
    finally:
        db.commit()
        if doc and doc.status in ["Procesado", "Error"]:
            notify_document_status(doc)
        db.close()


@app.post("/api/documents/upload")
async def upload_document(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...), 
    folder_id: Optional[int] = Form(None),
    vacancy_id: Optional[int] = Form(None),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role(["admin", "user", "service"]))
):
    ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
    MAX_FILE_SIZE = 10 * 1024 * 1024 # 10 MB

    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(status_code=400, detail="Formato de archivo no soportado. Solo .pdf, .docx y .txt permitidos.")

    file.file.seek(0, 2)
    size = file.file.tell()
    file.file.seek(0)
    if size > MAX_FILE_SIZE:
        raise HTTPException(status_code=400, detail="El archivo supera el límite de 10 MB")

    # Sanitización de nombre de archivo para prevenir Path Traversal
    raw_basename = os.path.basename(file.filename)
    safe_filename = re.sub(r'[^a-zA-Z0-9_\-\. ]', '_', raw_basename).strip()
    if not safe_filename or safe_filename.startswith("."):
        safe_filename = f"doc_{int(time.time())}{ext}"

    # Aislamiento de almacenamiento físico particionado por inquilino
    tenant_storage = f"storage/tenant_{current_user.tenant_id}"
    os.makedirs(tenant_storage, exist_ok=True)
    file_path = f"{tenant_storage}/{safe_filename}"
    with open(file_path, "wb") as buffer:
        shutil.copyfileobj(file.file, buffer)
        
    # Validación estricta multi-tenant de carpeta y vacante
    if folder_id:
        folder_check = db.query(Folder).filter(Folder.id == folder_id, Folder.tenant_id == current_user.tenant_id).first()
        if not folder_check:
            raise HTTPException(status_code=400, detail="La carpeta seleccionada no pertenece a tu organización")
    else:
        # Documentos recibidos por el bot de n8n van a "Bandeja n8n"
        if current_user.email == "n8n.service@gmail.com":
            folder = db.query(Folder).filter(Folder.name == "Bandeja n8n", Folder.tenant_id == current_user.tenant_id).first()
            if not folder:
                folder = Folder(name="Bandeja n8n", tenant_id=current_user.tenant_id)
                db.add(folder)
                db.commit()
                db.refresh(folder)
            folder_id = folder.id
        else:
            folder_id = None

    if vacancy_id:
        vac_check = db.query(Vacancy).filter(Vacancy.id == vacancy_id, Vacancy.tenant_id == current_user.tenant_id).first()
        if not vac_check:
            raise HTTPException(status_code=400, detail="La vacante seleccionada no pertenece a tu organización")
        
    doc = Document(
        tenant_id=current_user.tenant_id,
        filename=safe_filename, 
        file_path=file_path, 
        mime_type=file.content_type, 
        file_size_bytes=os.path.getsize(file_path), 
        status="Pendiente",
        folder_id=folder_id,
        vacancy_id=vacancy_id
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    background_tasks.add_task(process_document_background, doc.id, file_path, safe_filename)
    
    return {"message": "Documento subido. El procesamiento ha iniciado en segundo plano.", "id": doc.id, "status": doc.status}

@app.get("/api/documents")
def list_documents(
    filename: Optional[str] = Query(None, description="Filtrar por nombre de archivo"),
    category: Optional[str] = Query(None, description="Filtrar por categoría"),
    status: Optional[str] = Query(None, description="Filtrar por estado"),
    start_date: Optional[str] = Query(None, description="Fecha de inicio (YYYY-MM-DD)"),
    end_date: Optional[str] = Query(None, description="Fecha de fin (YYYY-MM-DD)"),
    db: Session = Depends(get_db), 
    current_user: User = Depends(get_current_user)
):
    query = db.query(Document).filter(Document.tenant_id == current_user.tenant_id)
    
    if filename:
        query = query.filter(Document.filename.ilike(f"%{filename}%"))
    if category:
        query = query.filter(Document.category == category)
    if status:
        query = query.filter(Document.status == status)
    if start_date:
        try:
            start = datetime.strptime(start_date, "%Y-%m-%d")
            query = query.filter(Document.created_at >= start)
        except ValueError:
            raise HTTPException(status_code=400, detail="Formato de start_date inválido. Use YYYY-MM-DD")
    if end_date:
        try:
            from datetime import timedelta
            end = datetime.strptime(end_date, "%Y-%m-%d") + timedelta(days=1)
            query = query.filter(Document.created_at < end)
        except ValueError:
            raise HTTPException(status_code=400, detail="Formato de end_date inválido. Use YYYY-MM-DD")
            
    docs = query.order_by(Document.id.desc()).all()
    user_role = getattr(current_user, "role", "user") or "user"
    return [serialize_document(d, user_role) for d in docs]

@app.delete("/api/documents/{doc_id}")
def delete_document(doc_id: int, db: Session = Depends(get_db), current_user: User = Depends(require_role(["admin"]))):
    doc = db.query(Document).filter(Document.id == doc_id, Document.tenant_id == current_user.tenant_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Documento no encontrado o no pertenece a tu organización")
    
    if os.path.exists(doc.file_path):
        try:
            os.remove(doc.file_path)
        except Exception:
            pass
        
    db.delete(doc)
    db.commit()
    return {"message": "Documento eliminado"}

@app.get("/api/documents/{doc_id}")
def get_document_details(doc_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    doc = db.query(Document).filter(Document.id == doc_id, Document.tenant_id == current_user.tenant_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Documento no encontrado o no pertenece a tu organización")
    user_role = getattr(current_user, "role", "user") or "user"
    return serialize_document(doc, user_role)

class DocumentUpdate(BaseModel):
    folder_id: Optional[int] = None
    folder_name: Optional[str] = None
    status: Optional[str] = None
    extracted_data: Optional[dict] = None

@app.patch("/api/documents/{doc_id}")
def update_document(
    doc_id: int, 
    update_data: DocumentUpdate, 
    db: Session = Depends(get_db), 
    current_user: User = Depends(require_role(["admin", "service"]))
):
    doc = db.query(Document).filter(Document.id == doc_id, Document.tenant_id == current_user.tenant_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Documento no encontrado o no pertenece a tu organización")
    
    if update_data.folder_name:
        folder = db.query(Folder).filter(Folder.name == update_data.folder_name, Folder.tenant_id == current_user.tenant_id).first()
        if not folder:
            folder = Folder(name=update_data.folder_name, tenant_id=current_user.tenant_id)
            db.add(folder)
            db.commit()
            db.refresh(folder)
        doc.folder_id = folder.id
    elif update_data.folder_id is not None:
        folder_check = db.query(Folder).filter(Folder.id == update_data.folder_id, Folder.tenant_id == current_user.tenant_id).first()
        if not folder_check:
            raise HTTPException(status_code=400, detail="La carpeta seleccionada no pertenece a tu organización")
        doc.folder_id = update_data.folder_id
        
    if update_data.status is not None:
        doc.status = update_data.status
    if update_data.extracted_data is not None:
        # Cifrar antes de guardar
        doc.extracted_data = {
            "_encrypted_data": encrypt_sensitive_data(update_data.extracted_data),
            "match_score": update_data.extracted_data.get("match_score"),
            "postulante": update_data.extracted_data.get("postulante"),
            "rol": update_data.extracted_data.get("rol"),
            "total": update_data.extracted_data.get("total"),
            "proveedor": update_data.extracted_data.get("proveedor")
        }
        
    db.commit()
    db.refresh(doc)
    return {"message": "Documento actualizado exitosamente", "id": doc.id, "folder_id": doc.folder_id, "status": doc.status}

class WorkflowUpdateSchema(BaseModel):
    target_folder_name: str
    payment_status: str
    notes: Optional[str] = None

@app.patch("/api/documents/{document_id}/workflow-update")
def update_document_from_n8n(
    document_id: int,
    payload: WorkflowUpdateSchema,
    db: Session = Depends(get_db),
    current_user: User = Depends(require_role(["admin", "service"]))
):
    from sqlalchemy.orm.attributes import flag_modified

    doc = db.query(Document).filter(
        Document.id == document_id,
        Document.tenant_id == current_user.tenant_id
    ).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Documento no encontrado o no pertenece a tu organización")

    folder = db.query(Folder).filter(
        Folder.name == payload.target_folder_name,
        Folder.tenant_id == current_user.tenant_id
    ).first()
    
    if not folder:
        folder = Folder(name=payload.target_folder_name, tenant_id=current_user.tenant_id)
        db.add(folder)
        db.commit()
        db.refresh(folder)

    doc.folder_id = folder.id

    raw_extracted = doc.extracted_data or {}
    if isinstance(raw_extracted, dict) and "_encrypted_data" in raw_extracted:
        raw_extracted = decrypt_sensitive_data(raw_extracted["_encrypted_data"])
    updated_data = dict(raw_extracted or {})
    updated_data["estado_contable"] = payload.payment_status
    if payload.notes:
        updated_data["notas_auditoria"] = payload.notes
        
    doc.extracted_data = {
        "_encrypted_data": encrypt_sensitive_data(updated_data),
        "match_score": updated_data.get("match_score"),
        "postulante": updated_data.get("postulante"),
        "rol": updated_data.get("rol"),
        "total": updated_data.get("total"),
        "proveedor": updated_data.get("proveedor"),
        "estado_contable": payload.payment_status
    }
    flag_modified(doc, "extracted_data")

    db.commit()
    return {
        "message": "Documento actualizado y movido con éxito",
        "document_id": doc.id,
        "folder": folder.name,
        "estado_contable": payload.payment_status
    }

@app.get("/api/documents/{doc_id}/download")
def download_document(doc_id: int, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    doc = db.query(Document).filter(Document.id == doc_id, Document.tenant_id == current_user.tenant_id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Documento no encontrado o no pertenece a tu organización")
    
    # Prevención estricta de Path Traversal
    real_path = os.path.abspath(doc.file_path)
    base_storage = os.path.abspath("storage")
    if not real_path.startswith(base_storage):
        raise HTTPException(status_code=403, detail="Acceso denegado: ruta de archivo no permitida")

    if not os.path.exists(doc.file_path):
        raise HTTPException(status_code=404, detail="Archivo no encontrado en el servidor")
        
    return FileResponse(path=doc.file_path, filename=doc.filename, media_type=doc.mime_type)

class RagQueryRequest(BaseModel):
    query: str
    document_ids: Optional[list[int]] = []

@app.post("/api/rag/query")
def rag_query(payload: RagQueryRequest, db: Session = Depends(get_db), current_user: User = Depends(get_current_user)):
    query = payload.query
    q_emb = genai.embed_content(model=MODEL_EMB, content=query, task_type="retrieval_query", output_dimensionality=768)['embedding']
    
    dist_expr = DocumentChunk.embedding.cosine_distance(q_emb)
    
    # Aislamiento multi-tenant estricto: solo chunks de documentos de este tenant
    base_query = db.query(DocumentChunk.chunk_text, Document.filename, dist_expr.label("dist")) \
                .join(Document, DocumentChunk.document_id == Document.id) \
                .filter(Document.tenant_id == current_user.tenant_id, dist_expr < 0.4)
                
    if payload.document_ids:
        # Validar que los document_ids pertenezcan exclusivamente a este tenant
        valid_ids = [r[0] for r in db.query(Document.id).filter(
            Document.id.in_(payload.document_ids), 
            Document.tenant_id == current_user.tenant_id
        ).all()]
        if valid_ids:
            base_query = base_query.filter(Document.id.in_(valid_ids))
        else:
            return {"answer": "Los documentos seleccionados no pertenecen a tu organización.", "sources": []}
        
    results = base_query.order_by(dist_expr).limit(4).all()
                
    if not results:
        return {"answer": "No encontré información relevante en los documentos de tu organización para responder a tu pregunta.", "sources": []}
        
    context = "\n".join([f"Archivo {r[1]}: {r[0]}" for r in results])
    
    model = genai.GenerativeModel(MODEL_TEXT)
    prompt = f"Responde la pregunta basándote SOLO en este contexto documental. Cita los archivos de soporte. Contexto:\n{context}\nPregunta: {query}"
    return {"answer": model.generate_content(prompt).text, "sources": list(set([r[1] for r in results]))}

